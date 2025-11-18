import os
import re
import zipfile
import rarfile
from PyPDF2 import PdfReader
from PIL import Image
import openpyxl
from docx import Document
import tempfile
from config import settings
import requests
import json


# OCR engine configuration
OCR_ENGINE = settings.OCR_ENGINE if hasattr(settings, 'OCR_ENGINE') else 'paddle'  # 'paddle' or 'tesseract'

# Initialize OCR engines (optional)
PADDLE_OCR_AVAILABLE = False
TESSERACT_OCR_AVAILABLE = False
paddle_ocr = None
tesseract_ocr = None

def initialize_ocr():
    """Initialize OCR engines if not already initialized"""
    global PADDLE_OCR_AVAILABLE, TESSERACT_OCR_AVAILABLE, paddle_ocr, tesseract_ocr
    
    # Initialize PaddleOCR if requested and not already initialized
    if OCR_ENGINE == 'paddle' and not PADDLE_OCR_AVAILABLE:
        try:
            from paddleocr import PaddleOCR
            paddle_ocr = PaddleOCR(lang='ch')
            PADDLE_OCR_AVAILABLE = True
        except ImportError:
            print("PaddleOCR not available - please install paddlepaddle and paddleocr packages")
            PADDLE_OCR_AVAILABLE = False
        except Exception as e:
            print(f"Error initializing PaddleOCR: {e}")
            PADDLE_OCR_AVAILABLE = False
    
    # Initialize Tesseract if requested and not already initialized
    if OCR_ENGINE == 'tesseract' and not TESSERACT_OCR_AVAILABLE:
        try:
            import pytesseract
            tesseract_ocr = pytesseract
            TESSERACT_OCR_AVAILABLE = True
        except ImportError:
            print("Tesseract not available - please install pytesseract package")
            TESSERACT_OCR_AVAILABLE = False
        except Exception as e:
            print(f"Error initializing Tesseract: {e}")
            TESSERACT_OCR_AVAILABLE = False


def extract_text_from_pdf(pdf_path):
    """Extract text from PDF file using PyPDF2"""
    try:
        text = ""
        reader = PdfReader(pdf_path)
        for page in reader.pages:
            text += page.extract_text() + "\n"
        return text
    except Exception as e:
        print(f"Error extracting text from PDF {pdf_path}: {str(e)}")
        return ""


def extract_text_from_docx(docx_path):
    """Extract text from DOCX file"""
    try:
        doc = Document(docx_path)
        text = "\n".join([paragraph.text for paragraph in doc.paragraphs])
        return text
    except Exception as e:
        print(f"Error extracting text from DOCX {docx_path}: {str(e)}")
        return ""


def extract_text_from_xlsx(xlsx_path):
    """Extract text from XLSX file"""
    try:
        workbook = openpyxl.load_workbook(xlsx_path)
        text = ""
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                row_text = " ".join([str(cell) if cell else "" for cell in row])
                text += row_text + "\n"
        return text
    except Exception as e:
        print(f"Error extracting text from XLSX {xlsx_path}: {str(e)}")
        return ""


def extract_text_from_txt(txt_path):
    """Extract text from TXT file"""
    try:
        with open(txt_path, 'r', encoding='utf-8') as file:
            return file.read()
    except Exception as e:
        print(f"Error extracting text from TXT {txt_path}: {str(e)}")
        return ""


def extract_ocr_from_image(image_path):
    """Extract text from image using configured OCR engine"""
    initialize_ocr()
    
    if OCR_ENGINE == 'paddle' and PADDLE_OCR_AVAILABLE:
        try:
            result = paddle_ocr.ocr(image_path, cls=True)
            text = ""
            for page_result in result:
                if page_result:  # Check if result is not None
                    for item in page_result:
                        if item and len(item) > 1:
                            text += item[1][0] + " "  # Get the recognized text
            return text
        except Exception as e:
            print(f"Error performing PaddleOCR on image {image_path}: {str(e)}")
            return ""
    elif OCR_ENGINE == 'tesseract' and TESSERACT_OCR_AVAILABLE:
        try:
            from PIL import Image
            img = Image.open(image_path)
            text = tesseract_ocr.image_to_string(img, lang='chi_sim+eng')
            return text
        except Exception as e:
            print(f"Error performing Tesseract OCR on image {image_path}: {str(e)}")
            return ""
    else:
        if OCR_ENGINE == 'paddle':
            return "PaddleOCR not available - please install paddlepaddle and paddleocr packages"
        elif OCR_ENGINE == 'tesseract':
            return "Tesseract not available - please install pytesseract package"
        else:
            return f"Unsupported OCR engine: {OCR_ENGINE}"


def extract_ocr_from_pdf(pdf_path):
    """Extract text from PDF using OCR (for image-based PDFs)"""
    initialize_ocr()
    
    if OCR_ENGINE == 'paddle' and not PADDLE_OCR_AVAILABLE:
        if OCR_ENGINE == 'paddle':
            return "PaddleOCR not available - please install paddlepaddle and paddleocr packages"
        elif OCR_ENGINE == 'tesseract':
            return "Tesseract not available - please install pytesseract package"
        else:
            return f"Unsupported OCR engine: {OCR_ENGINE}"
    
    try:
        # For image-based PDFs, we need to convert each page to an image first
        import fitz  # PyMuPDF

        doc = fitz.open(pdf_path)
        all_text = ""

        for page_num in range(len(doc)):
            page = doc.load_page(page_num)
            pix = page.get_pixmap()

            # Save as temporary image
            with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp_file:
                pix.save(tmp_file.name)
                text = extract_ocr_from_image(tmp_file.name)
                all_text += text + "\n"

                # Clean up temporary file
                os.unlink(tmp_file.name)

        doc.close()
        return all_text
    except Exception as e:
        print(f"Error performing OCR on PDF {pdf_path}: {str(e)}")
        return ""


def extract_text_from_file(file_path):
    """Extract text from various file types"""
    _, ext = os.path.splitext(file_path.lower())

    if ext == '.pdf':
        # Try to extract text first, if it fails or returns little text, use OCR
        text = extract_text_from_pdf(file_path)
        if len(text.strip()) < 100:  # If less than 100 characters, try OCR
            ocr_text = extract_ocr_from_pdf(file_path)
            return max(text, ocr_text, key=len)  # Return the longer text
        return text
    elif ext in ['.docx']:
        return extract_text_from_docx(file_path)
    elif ext in ['.xlsx', '.xls']:
        return extract_text_from_xlsx(file_path)
    elif ext in ['.txt']:
        return extract_text_from_txt(file_path)
    elif ext in ['.jpg', '.jpeg', '.png', '.bmp', '.gif', '.tiff']:
        return extract_ocr_from_image(file_path)
    else:
        return ""


def extract_zip_content(zip_path, extract_to):
    """Extract content from zip/rar files"""
    try:
        os.makedirs(extract_to, exist_ok=True)
        if zip_path.endswith('.zip'):
            with zipfile.ZipFile(zip_path, 'r') as zip_ref:
                zip_ref.extractall(extract_to)
        elif zip_path.endswith('.rar'):
            with rarfile.RarFile(zip_path, 'r') as rar_ref:
                rar_ref.extractall(extract_to)
    except Exception as e:
        print(f"Error extracting archive {zip_path}: {str(e)}")


def contains_id_card(text):
    """Check if text contains ID card numbers"""
    if not text:
        return False

    # Pattern for Chinese ID card number (18 digits, with possible X at the end)
    pattern = r'\b[1-9]\d{5}(18|19|20)\d{2}((0[1-9])|(1[0-2]))(([0-2][1-9])|10|20|30|31)\d{3}[0-9Xx]\b'
    matches = re.findall(pattern, text)
    return len(matches) > 0


def contains_phone(text):
    """Check if text contains phone numbers"""
    if not text:
        return False

    # Pattern for Chinese phone numbers
    pattern = r'(\b(?:\+?86[-\s]?)?(?:1[3-9]\d{9}|(?:[0-9]{3,4}[-\s]?)?[0-9]{7,8})\b)'
    matches = re.findall(pattern, text)
    return len(matches) > 0


def get_llm_content(image_path):
    """Get content from image using LLM (OpenAI GPT-4 Vision or similar)"""
    api_key = settings.OPENAI_API_KEY
    api_base = settings.OPENAI_BASE_URL
    model = settings.MODEL
    
    if not api_key:
        return "OpenAI API key not configured"

    try:
        # Read image and encode to base64
        import base64
        with open(image_path, "rb") as image_file:
            encoded_image = base64.b64encode(image_file.read()).decode('ascii')

        headers = {
            "Content-Type": "application/json",
            "Authorization": f"Bearer {api_key}"
        }

        payload = {
            "model": model,
            "messages": [
                {
                    "role": "user",
                    "content": [
                        {
                            "type": "text",
                            "text": "Describe the content of this image in detail, focusing on any text, documents, or important information visible in the image."
                        },
                        {
                            "type": "image_url",
                            "image_url": {
                                "url": f"data:image/jpeg;base64,{encoded_image}"
                            }
                        }
                    ]
                }
            ],
            "max_tokens": 500
        }

        response = requests.post(api_base + "/chat/completions", headers=headers, json=payload)
        response.raise_for_status()

        result = response.json()
        return result['choices'][0]['message']['content']
    except Exception as e:
        print(f"Error getting LLM content for {image_path}: {str(e)}")
        return ""


def get_content_analysis(content):
    """Get content analysis from LLM to identify sensitive information"""
    api_key = settings.OPENAI_API_KEY
    api_base = settings.OPENAI_BASE_URL
    model = settings.MODEL
    prompt = settings.PROMPTS
    
    if not api_key:
        return "OpenAI API key not configured"

    try:
        headers = {
            "Content-Type": "application/json",
            "Authorization": f"Bearer {api_key}"
        }

        payload = {
            "model": model,
            "messages": [
                {
                    "role": "user",
                    "content": f"{prompt}\n\nContent to analyze:\n{content}"
                }
            ],
            "max_tokens": 500
        }

        response = requests.post(api_base + "/chat/completions", headers=headers, json=payload)
        response.raise_for_status()

        result = response.json()
        return result['choices'][0]['message']['content']
    except Exception as e:
        print(f"Error getting content analysis: {str(e)}")
        return ""


def detect_sensitive_info_ai(content):
    """Detect sensitive information using AI analysis"""
    analysis = get_content_analysis(content)
    has_id_card = contains_id_card(content) or "id card" in analysis.lower() or "identity card" in analysis.lower()
    has_phone = contains_phone(content) or "phone" in analysis.lower() or "mobile" in analysis.lower()
    return has_id_card, has_phone, analysis